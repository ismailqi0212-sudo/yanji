import asyncio
import json
import pytest
from yanji.common import AppError, clone
from yanji.agent.cancellation import CancellationToken, Cancelled
from yanji.agent.loop import parse_decision
from yanji.agent.runtime_factory import AgentRuntimeFactory
from yanji.context.assembler import ContextAssembler
from yanji.llm.model import ModelService
from yanji.memory.store import Store
from yanji.trace.session_store import SessionStore
from yanji.tools.registry import ToolRegistry
from yanji.tools.execution import ToolExecutionPipeline
from yanji.research.artifacts import make_card
from yanji.research.validation import validate_answer
from yanji.research.metrics import align_operating_facts
from yanji.tools.document_analysis import (
    validate_operating_facts,
    validate_part,
    validate_series,
    batches,
)
from yanji.verification.verifier import VerificationRejected
from conftest import call


def session(tmp_path):
    store = Store(tmp_path)
    sessions = SessionStore(store)
    chat = store.create()
    run = dict(
        id="run",
        messageId="answer",
        question="研究",
        status="running",
        sourceMode="external",
        mode="research",
        route={"maxFindings": 5},
        records=[],
        events=[],
        usage=0,
    )
    sessions.create_run(
        chat["id"],
        run,
        dict(id="question", role="user", text="研究", cardIds=[]),
        dict(id="answer", role="assistant", text="", cardIds=[]),
    )
    return store, sessions, chat["id"], run


async def test_model_ledger_parent_scope_concurrency_and_exactly_once(tmp_path):
    store, sessions, cid, run = session(tmp_path)

    async def transport(*args, **kwargs):
        return dict(content="ok", usage=7)

    model = ModelService(
        transport, sessions, dict(sessionId=cid, runId=run["id"], purpose="decision")
    )
    child = model.scope(parentCallId="tool", purpose="extract")
    await asyncio.gather(
        child.complete([], metadata={"batch": 1}),
        child.complete([], metadata={"batch": 2}),
        model.complete([]),
    )
    trace = sessions.read_trace(cid, run["id"])
    assert len(trace["modelCalls"]) == 3
    assert sessions.load(cid, run["id"])["usage"] == 21
    sessions.finish_model_call(trace["modelCalls"][0])
    assert sessions.load(cid, run["id"])["usage"] == 21
    assert {c["batch"] for c in trace["modelCalls"] if c.get("parentCallId")} == {1, 2}
    store.db.close()


async def test_model_cancel_records_unknown_usage(tmp_path):
    store, sessions, cid, run = session(tmp_path)
    token = CancellationToken()

    async def transport(*args, **kwargs):
        await asyncio.sleep(100)

    model = ModelService(
        transport, sessions, dict(sessionId=cid, runId=run["id"]), token
    )
    task = asyncio.create_task(model.complete([]))
    await asyncio.sleep(0.01)
    token.cancel()
    with pytest.raises(Cancelled):
        await task
    saved = sessions.read_trace(cid, run["id"])["modelCalls"][0]
    assert saved["status"] == "cancelled" and saved["usageKnown"] is False
    store.db.close()


async def test_registry_lazy_loading_and_source_category_permissions():
    loaded = []

    def load():
        loaded.append(1)

        async def handler(args, context):
            return {"value": args["value"]}

        return handler

    entry = dict(
        name="test",
        description="test",
        parameters={
            "type": "object",
            "properties": {"value": {"type": "number"}},
            "required": ["value"],
        },
        category="data",
        sourceModes=["external"],
        load=load,
    )
    denied = ToolRegistry([entry], source_mode="documents")
    assert denied.list() == []
    with pytest.raises(AppError):
        await denied.invoke("test", {"value": 1})
    assert loaded == []
    registry = ToolRegistry([entry], source_mode="external", categories=["data"])
    assert loaded == []
    assert (await registry.invoke("test", {"value": 1}))["value"] == 1
    assert (await registry.invoke("test", {"value": 2}))["value"] == 2
    assert loaded == [1]
    with pytest.raises(AppError):
        await registry.invoke("test", {})


async def test_verifier_rejection_and_commit_conflict_leave_no_artifact(tmp_path):
    store, sessions, cid, run = session(tmp_path)
    token = CancellationToken()

    async def transport(*args, **kwargs):
        return {"usage": 1}

    model = ModelService(transport, sessions, dict(sessionId=cid, runId="run"), token)
    ctx = dict(
        cid=cid,
        run=run,
        sessions=sessions,
        token=token,
        seen=set(),
        calculations=[],
        financial_cache=set(),
        model=model,
    )

    async def handler(args, context):
        source = context["workspace"].source({"value": 1}, "候选来源")
        context["workspace"].artifact(
            make_card("brief", "候选", {"body": "内容"}, [source["id"]])
        )
        context["seen"].add(source["id"])
        return {"sourceId": source["id"]}

    class Reject:
        async def verify(self, *args):
            return {"accepted": False, "reason": "拒绝"}

    pipeline = ToolExecutionPipeline(ctx, Reject())
    with pytest.raises(VerificationRejected):
        await pipeline.run({"name": "test"}, {}, {"callId": "rejected"}, handler)
    assert (
        store.get(cid)["cards"] == []
        and store.get(cid)["researchSources"] == []
        and ctx["seen"] == set()
    )

    class ConcurrentEdit:
        async def verify(self, *args):
            chat = store.get(cid)
            chat["cards"].append(make_card("brief", "用户编辑", {}))
            store.put(chat)
            return {"accepted": True}

    with pytest.raises(AppError, match="研究内容已变化"):
        await ToolExecutionPipeline(ctx, ConcurrentEdit()).run(
            {"name": "test"}, {}, {"callId": "conflict"}, handler
        )
    assert [c["title"] for c in store.get(cid)["cards"]] == ["用户编辑"]
    assert store.get(cid)["researchSources"] == []
    store.db.close()


async def test_cancellation_during_verification_prevents_commit(tmp_path):
    store, sessions, cid, run = session(tmp_path)
    token = CancellationToken()

    async def transport(*args, **kwargs):
        return {"usage": 1}

    ctx = dict(
        cid=cid,
        run=run,
        sessions=sessions,
        token=token,
        seen=set(),
        calculations=[],
        financial_cache=set(),
        model=ModelService(
            transport, sessions, dict(sessionId=cid, runId="run"), token
        ),
    )

    class Stop:
        async def verify(self, *args):
            token.cancel()
            return {"accepted": True}

    async def handler(args, c):
        c["workspace"].artifact(make_card("brief", "不可提交", {}))
        return {}

    with pytest.raises(Cancelled):
        await ToolExecutionPipeline(ctx, Stop()).run(
            {"name": "test"}, {}, {"callId": "stop"}, handler
        )
    assert store.get(cid)["cards"] == []
    store.db.close()


async def test_trace_commit_failure_rolls_back_business_write(tmp_path):
    store, sessions, cid, run = session(tmp_path)
    base = store.get(cid)
    candidate = clone(base)
    candidate["cards"].append(make_card("brief", "候选", {}))

    def fail(*args):
        raise RuntimeError("audit failed")

    sessions.audit = fail
    with pytest.raises(RuntimeError):
        sessions.commit_tool(cid, run["id"], base, candidate, {"name": "test"})
    assert store.get(cid)["cards"] == []
    store.db.close()


async def test_runtime_children_are_disabled_or_scope_restricted():
    async def assemble(*args):
        return {}

    factory = AgentRuntimeFactory(assemble)
    runtime = factory.create("c", categories=["analysis"], source_mode="documents")
    with pytest.raises(AppError):
        runtime.create_child()
    factory = AgentRuntimeFactory(assemble, allow_children=True)
    runtime = factory.create(
        "c", categories=["analysis"], source_mode="documents", max_steps=3
    )
    with pytest.raises(AppError):
        runtime.create_child(categories=["data"])
    with pytest.raises(AppError):
        runtime.create_child(source_mode="external")
    with pytest.raises(AppError):
        runtime.create_child(max_steps=4)
    child = runtime.create_child(categories=["analysis"], max_steps=2)
    assert child.scope["parentRunId"] == runtime.scope["runId"]
    runtime.token.cancel()
    assert child.token.cancelled


async def test_context_rebuild_is_pure_and_filters_long_term_sources():
    class Registry:
        def list(self):
            return [
                dict(
                    name="done",
                    description="done",
                    parameters={"type": "object"},
                    terminal=True,
                )
            ]

    class Short:
        async def retrieve(self, state):
            return dict(history=[], projectMemory={}, documentSourceIds=["D1"])

    class Long:
        async def retrieve(self, scope):
            return dict(
                sources=[
                    dict(id="D1", kind="document"),
                    dict(id="S1", kind="external"),
                ],
                items=[
                    dict(title="allowed", sourceIds=["D1"]),
                    dict(title="forbidden", sourceIds=["S1"]),
                ],
            )

    assembler = ContextAssembler(
        Registry(),
        short_term=Short(),
        long_term=Long(),
        perspectives=[{"instruction": "保持看空反证视角"}],
    )
    state = dict(
        task={"sourceMode": "documents", "goal": "研究"},
        trajectory=[],
        progress={"toolCalls": 0},
        limits={"maxToolCalls": 3},
    )
    before = clone(state)
    output = await assembler.build(state)
    assert state == before
    text = json.dumps(output, ensure_ascii=False)
    assert "allowed" in text and "forbidden" not in text
    assert "保持看空反证视角" in text


@pytest.mark.parametrize(
    "response,kind",
    [
        (dict(content="完成"), "invalid"),
        (dict(content='{"type":"finish","summary":"完成"}'), "finish"),
        (call("test", {}), "act"),
    ],
)
def test_decision_protocol(response, kind):
    assert parse_decision(response)["kind"] == kind


@pytest.mark.parametrize(
    "kind,metric", [("flow", "付费用户"), ("flow", "ARPU"), ("stock", "营业收入")]
)
def test_nonadditive_metrics_never_derive_quarters(kind, metric):
    facts = [
        dict(
            company="Example",
            metric=metric,
            period=p,
            value=v,
            rawValue=str(v),
            unit="元",
            currency="CNY",
            scope="合并",
            kind=kind,
            calendar="12月",
            sourceId="D1",
        )
        for p, v in [("2024-Q1", 10), ("2024-H1", 30)]
    ]
    result = align_operating_facts(facts)
    assert "2024-Q2" not in result["rows"][0]["cells"]


def test_extraction_rejects_invented_quote_number_and_duplicate_chart_periods():
    chunks = [
        dict(
            id="D1",
            text="2024年度收入120万元，相关财务披露需要核对。",
            page=1,
            documentId="doc",
        )
    ]
    fact = dict(
        company="E",
        metric="收入",
        period="2024",
        kind="flow",
        value=999,
        rawValue="999",
        unit="万元",
        currency="CNY",
        scope="合并",
        calendar="12月",
        sourceId="D1",
        quote=chunks[0]["text"],
    )
    assert validate_operating_facts({"facts": [fact]}, chunks)["facts"] == []
    assert (
        validate_part(
            {"findings": [{"sourceId": "D1", "quote": "不存在的原文引句内容"}]}, chunks
        )["findings"]
        == []
    )
    with pytest.raises(AppError):
        validate_series(
            dict(
                title="收入",
                unit="万元",
                scope="年度",
                rows=[dict(period="2024", value=120, rawValue="120", sourceId="D1")]
                * 2,
            ),
            chunks,
        )
    assert sum(map(len, batches([dict(text="x" * 2000) for _ in range(4)], 3000))) == 4


def test_evidence_scope_cannot_borrow_unrelated_calculation():
    chat = dict(
        documents=[],
        researchSources=[dict(id="S1", type="web", text="{}", data={})],
        cards=[],
    )
    with pytest.raises(AppError, match="未匹配"):
        validate_answer(
            chat,
            dict(answer="增长20%[来源](evidence:S1)", sourceIds=["S1"]),
            {"S1"},
            [dict(operation="growth", result=20, sourceIds=["S2"])],
        )


def test_node_ciphertext_format_is_compatible(tmp_path):
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    store = Store(tmp_path)
    iv = bytes(range(12))
    key = "fake-node-format-value"
    encrypted = AESGCM(store.key_path.read_bytes()).encrypt(iv, key.encode(), None)
    settings = {
        "baseUrl": "https://example.test",
        "model": "test",
        "key": ".".join(x.hex() for x in (iv, encrypted[:-16], encrypted[-16:])),
    }
    store.db.execute("INSERT INTO settings VALUES(1,?)", (json.dumps(settings),))
    assert store.settings()["key"] == key
    assert key not in json.dumps(store.public_settings())
    store.db.close()


async def test_lazy_async_loader_runs_once_under_concurrent_invocations():
    loads = []

    async def loader():
        loads.append(1)
        await asyncio.sleep(0.01)

        async def handler(args, context):
            return args

        return handler

    entry = dict(
        name="test", description="test", parameters={"type": "object"}, load=loader
    )
    registry = ToolRegistry([entry])
    assert await asyncio.gather(
        registry.invoke("test", {"n": 1}), registry.invoke("test", {"n": 2})
    ) == [{"n": 1}, {"n": 2}]
    assert loads == [1]


async def test_report_selection_conflict_prevents_stale_commit(env):
    from conftest import wait_run

    app, client, model = env
    cid = (await client.post("/api/chats", json={})).json()["id"]
    run = (
        await client.post(f"/api/chats/{cid}/messages", json={"text": "比较财务"})
    ).json()
    await wait_run(app, cid, run["id"])
    card = next(c for c in app.state.store.get(cid)["cards"] if c["kind"] == "brief")
    await client.patch(f"/api/chats/{cid}/cards/{card['id']}", json={"selected": True})
    transport = app.state.reports.transport

    async def changed_selection(*args, **kwargs):
        chat = app.state.store.get(cid)
        next(c for c in chat["cards"] if c["id"] == card["id"])["selected"] = False
        app.state.store.put(chat)
        return await transport(*args, **kwargs)

    app.state.reports.transport = changed_selection
    response = await client.post(f"/api/chats/{cid}/compile", json={})
    assert response.status_code == 409, response.text
    assert app.state.store.get(cid)["report"] is None


async def test_report_can_be_cancelled_via_existing_stop_endpoint(env):
    from conftest import wait_run

    app, client, model = env
    cid = (await client.post("/api/chats", json={})).json()["id"]
    run = (
        await client.post(f"/api/chats/{cid}/messages", json={"text": "比较财务"})
    ).json()
    await wait_run(app, cid, run["id"])
    card = next(c for c in app.state.store.get(cid)["cards"] if c["kind"] == "brief")
    await client.patch(f"/api/chats/{cid}/cards/{card['id']}", json={"selected": True})
    model.wait = asyncio.Event()
    task = asyncio.create_task(client.post(f"/api/chats/{cid}/compile", json={}))
    await asyncio.sleep(0.02)
    rid = app.state.store.get(cid)["reportJob"]["runId"]
    await client.post(f"/api/chats/{cid}/runs/{rid}/stop", json={})
    response = await task
    assert response.status_code == 400
    chat = app.state.store.get(cid)
    assert chat["report"] is None and chat["reportJob"]["status"] == "cancelled"
    assert (
        app.state.sessions.read_trace(cid, rid)["modelCalls"][0]["status"]
        == "cancelled"
    )


def test_docx_tables_are_not_lost_during_python_parsing():
    import io
    from docx import Document
    from yanji.tools.documents import parse_document

    document = Document()
    document.add_paragraph("公司经营情况及收入披露，请核对以下经营数据。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "营业收入"
    table.cell(0, 1).text = "2024"
    table.cell(1, 0).text = "订阅服务"
    table.cell(1, 1).text = "120万元"
    stream = io.BytesIO()
    document.save(stream)
    parsed = parse_document(stream.getvalue(), "report.docx")
    text = "\n".join(p["text"] for p in parsed["pages"])
    assert "订阅服务" in text and "120万元" in text


def test_legacy_inherited_configuration_uses_original_cipher_layout(
    tmp_path, monkeypatch
):
    import os, sqlite3
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    monkeypatch.delenv("CHAT_NO_INHERIT", raising=False)
    old = tmp_path / "thesis-ledger/data"
    old.mkdir(parents=True)
    key = os.urandom(32)
    (old / ".credential-key").write_bytes(key)
    iv = os.urandom(12)
    encrypted = AESGCM(key).encrypt(iv, b"fake-inherited-key", None)
    conn = sqlite3.connect(old / "ledger.sqlite")
    conn.execute("CREATE TABLE settings(id INTEGER PRIMARY KEY,body TEXT)")
    conn.execute(
        "INSERT INTO settings VALUES(1,?)",
        (
            json.dumps(
                {
                    "baseUrl": "https://example.test",
                    "model": "test",
                    "key": ".".join(
                        x.hex() for x in (iv, encrypted[-16:], encrypted[:-16])
                    ),
                }
            ),
        ),
    )
    conn.commit()
    conn.close()
    store = Store(tmp_path / "yanji-chat/data")
    monkeypatch.delenv("YANJI_IMPORT_LEGACY_SETTINGS", raising=False)
    assert store.settings()["inherited"] is False
    assert store.settings()["key"] == ""
    monkeypatch.setenv("YANJI_IMPORT_LEGACY_SETTINGS", "1")
    assert store.settings()["inherited"] is True
    assert store.settings()["key"] == "fake-inherited-key"
    store.db.close()
